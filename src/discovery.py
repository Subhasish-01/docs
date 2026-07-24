from docx import Document


class DocumentDiscovery:

    def __init__(self, document_path):
        self.document_path = document_path
        self.document = None

    def load_document(self):
       try:
            self.document = Document(self.document_path)
            print("✓ Document loaded successfully")
       except Exception as e:
            print(f"Error loading document: {e}")
            raise

    def extract_paragraphs(self):
        """
        Extract all paragraphs from the document.
        """
        return self.document.paragraphs

    def inspect_runs(self, paragraph):
        print("\nParagraph:")
        print(paragraph.text)

        print("\nRuns:")

        for i, run in enumerate(paragraph.runs, start=1):
            print(f"Run {i}")
            print(f"Text   : {run.text}")
            print(f"Bold   : {run.bold}")
            print(f"Italic : {run.italic}")
            print("-" * 30)

    def extract_text(self):
        """
        Extract all non-empty text from the document.
        Returns a single string.
        """

        full_text = []

        for paragraph in self.document.paragraphs:
            if paragraph.text.strip():
                full_text.append(paragraph.text)

        return "\n".join(full_text)